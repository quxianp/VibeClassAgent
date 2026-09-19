"""文档生成：Word（python-docx）+ PDF（fpdf2 直出）。

设计要点：
  - **不再依赖 LibreOffice**：PDF 由 fpdf2 直接生成，版式可控、速度快、部署轻；
  - 中文字体显式指定（Word 用微软雅黑，PDF 用黑体 TTF），避免缺字；
  - 截图数量受 `max_screenshots` 限制，防止文档体积失控；
  - 两个格式**互相独立**：任一个失败不影响另一个，返回值里分别报告。
"""

from __future__ import annotations

import json
import os
from typing import Any, Dict, List, Optional


def _safe_name(s: str) -> str:
    """把课程名等转成安全文件名。"""
    bad = '<>:"/\\|?*'
    out = "".join("_" if c in bad else c for c in s).strip()
    return out or "未命名"


def _load_summary(path: str) -> Dict[str, Any]:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _load_transcript(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except Exception:
        return ""


def _set_base_font(doc) -> None:
    """把 Word 文档默认字体设为中文友好字体。"""
    try:
        from docx.oxml.ns import qn
        from docx.shared import Pt

        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass


def build_docx(
    course: str,
    date: str,
    start: str,
    end: str,
    summary: Dict[str, Any],
    transcript: str,
    screenshots: List[str],
    out_dir: str,
    max_screenshots: int = 6,
    include_transcript: bool = True,
) -> Optional[str]:
    """生成 Word 文档；缺 python-docx 时抛异常。"""
    try:
        from docx import Document  # type: ignore
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError(
            "未安装 python-docx。请执行：python -m pip install -r requirements.txt"
        ) from exc

    os.makedirs(out_dir, exist_ok=True)
    doc = Document()
    _set_base_font(doc)

    title = summary.get("title") or course
    doc.add_heading(f"{course} 课堂纪要", level=0)
    doc.add_paragraph(f"日期：{date}　时间：{start} - {end}")
    if title and title != course:
        doc.add_paragraph(f"主题：{title}")

    points = summary.get("points") or []
    if points:
        doc.add_heading("课堂重点", level=1)
        for p in points:
            doc.add_paragraph(str(p), style="List Number")

    notices = summary.get("notices") or []
    if notices:
        doc.add_heading("重点通知", level=1)
        for n in notices:
            doc.add_paragraph(str(n), style="List Bullet")

    keywords = summary.get("keywords") or []
    if keywords:
        doc.add_heading("关键词", level=1)
        doc.add_paragraph("、".join(str(k) for k in keywords))

    shots = [s for s in screenshots if s and os.path.exists(s)][:max_screenshots]
    if shots:
        doc.add_heading("关键截图", level=1)
        for s in shots:
            try:
                doc.add_picture(s, width=Inches(6.0))
            except Exception:
                continue
            doc.add_paragraph(os.path.basename(s))

    if include_transcript and transcript.strip():
        doc.add_heading("附：转写全文", level=1)
        for line in transcript.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())

    path = os.path.join(out_dir, f"{_safe_name(course)}_{date.replace('-', '')}_课堂纪要.docx")
    doc.save(path)
    return path


def generate(params: Dict[str, Any]) -> Dict[str, Any]:
    """入口：按参数生成 Word / PDF。

    参数（来自 Rust 侧 pipeline）：
      course, date, start, end, teacher_id,
      summary_path, transcript_path, screenshots[], out_dir, formats[]
    """
    course = str(params.get("course") or "未命名课程")
    date = str(params.get("date") or "")
    start = str(params.get("start") or "")
    end = str(params.get("end") or "")
    out_dir = str(params.get("out_dir") or ".")
    formats = [str(f).lower() for f in (params.get("formats") or ["docx"])]
    screenshots = list(params.get("screenshots") or [])
    max_shots = int(params.get("max_screenshots") or 6)

    summary = _load_summary(str(params.get("summary_path") or ""))
    transcript = _load_transcript(str(params.get("transcript_path") or ""))

    docx_path = None
    pdf_path = None
    notes: List[str] = []

    if "docx" in formats:
        try:
            docx_path = build_docx(
                course, date, start, end, summary, transcript,
                screenshots, out_dir, max_shots,
            )
        except Exception as exc:  # noqa: BLE001
            notes.append(f"Word 生成失败：{str(exc)[:120]}")

    if "pdf" in formats:
        try:
            from .pdfgen import build_pdf

            pdf_path = build_pdf(
                course, date, start, end, summary, transcript,
                screenshots, out_dir, max_shots,
            )
        except Exception as exc:  # noqa: BLE001
            notes.append(f"PDF 生成失败：{str(exc)[:120]}")

    return {
        "docx": docx_path or "",
        "pdf": pdf_path or "",
        "notes": notes,
    }
